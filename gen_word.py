import json
import random
import uuid
import os
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

# 定义数据和模板目录
DATA_DIR = "data"

def load_data(q_type: str = "抗抑郁型") -> dict[str, str | list]:
    """加载测试数据，根据问卷类型动态加载。"""
    os.makedirs(DATA_DIR, exist_ok=True)
    
    # 1. 加载首尾提示语（如果存在该类型的专属提示语则使用，否则用默认的）
    paper_file = os.path.join(DATA_DIR, f"{q_type}_paper.txt")
    if not os.path.exists(paper_file):
        paper_file = "paper_example.txt" # 默认抑郁型首尾语
        
    with open(paper_file, "r", encoding="utf-8") as rf:
        paper_hint = rf.read()
    
    if "---\n" not in paper_hint:
        raise ValueError(f"文件 {paper_file} 中不包含 `---\\n` 分隔符")
    head_hint, finish_hint = paper_hint.split("---\n")

    # 2. 加载题库（如果该类型题库存在则读取，否则读取默认 titles.json）
    question_file = os.path.join(DATA_DIR, f"{q_type}.json")
    if not os.path.exists(question_file):
        question_file = "titles.json" # 默认抑郁型题库
        
    with open(question_file, "r", encoding="utf-8") as jf:
        questions = json.load(jf)

    return {
        "head": head_hint,
        "end": finish_hint,
        "problems": questions
    }

def generate_word(q_type: str = "抗抑郁型") -> str:
    """生成一份随机的心理调查问卷 Word 文档。"""
    doc = Document()
    doc.add_heading(f"{q_type}调查问卷", level=1)

    data = load_data(q_type)
    doc.add_paragraph(data["head"])

    problems = data["problems"]
    os.makedirs("doc", exist_ok=True)

    question_num = 1
    # 去掉原来的 range(5) 限制，动态遍历所有组，防止非抑郁类题目组数不同报错
    for group in problems:
        num_to_pick = min(2, len(group))
        selected_problems = random.sample(group, num_to_pick)

        for problem in selected_problems:
            text = f"{question_num}. "
            if problem.get("is_reverse_scored", False):
                text += "【本题需要反向评分】"
            text += problem["question_text"]
            doc.add_paragraph(text)
            question_num += 1

    doc.add_paragraph(data["end"])
    
    # 修改字体
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(12)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), "宋体")

    file_name = str(uuid.uuid4())[:8] + ".docx"
    file_path = os.path.join("doc", file_name)
    doc.save(file_path)

    return file_name